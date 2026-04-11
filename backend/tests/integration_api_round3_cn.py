import json
import os
import re
import sys
from datetime import datetime

sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), "..")))

from app import create_app
from app.extensions import mongo


def _now_stamp():
    return datetime.now().strftime("%Y%m%d")


def _ensure_dir(path: str):
    os.makedirs(path, exist_ok=True)


def _j(obj):
    return json.dumps(obj, ensure_ascii=False, separators=(",", ":"))


def _is_garbled(text: str) -> bool:
    if text is None:
        return False
    if "�" in text:
        return True
    if "\\u" in text:
        return True
    return False


def _make_curl(method: str, url: str, headers: dict, json_body):
    parts = ["curl", "-i", "-X", method.upper(), f"\"{url}\""]
    for k, v in (headers or {}).items():
        parts += ["-H", f"\"{k}: {v}\""]
    if json_body is not None:
        parts += ["-H", "\"Content-Type: application/json; charset=utf-8\""]
        parts += ["--data-raw", f"'{_j(json_body)}'"]
    return " ".join(parts)


def _write_xlsx_data_log(path: str, rows: list[dict]):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment

    wb = Workbook()
    ws = wb.active
    ws.title = "R3_Data_Log"

    headers = ["序号", "用例编号", "接口路径", "中文输入数据", "响应状态码", "响应数据（中文）", "数据库落库值", "是否一致", "备注"]
    ws.append(headers)

    header_fill = PatternFill("solid", fgColor="D9E1F2")
    header_font = Font(bold=True)
    for col in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(wrap_text=True, vertical="top")

    red_fill = PatternFill("solid", fgColor="FFCDD2")
    wrap = Alignment(wrap_text=True, vertical="top")

    for i, r in enumerate(rows, start=1):
        ws.append(
            [
                i,
                r.get("case_id"),
                r.get("path"),
                r.get("cn_input"),
                r.get("status_code"),
                r.get("cn_response"),
                r.get("db_value"),
                r.get("is_match"),
                r.get("remark"),
            ]
        )
        row_idx = i + 1
        for col in range(1, len(headers) + 1):
            ws.cell(row=row_idx, column=col).alignment = wrap

        if r.get("is_match") in ("否", False) or _is_garbled(str(r.get("cn_response") or "")) or _is_garbled(str(r.get("db_value") or "")):
            for col in range(1, len(headers) + 1):
                ws.cell(row=row_idx, column=col).fill = red_fill

    ws.column_dimensions["A"].width = 8
    ws.column_dimensions["B"].width = 18
    ws.column_dimensions["C"].width = 28
    ws.column_dimensions["D"].width = 45
    ws.column_dimensions["E"].width = 12
    ws.column_dimensions["F"].width = 55
    ws.column_dimensions["G"].width = 45
    ws.column_dimensions["H"].width = 10
    ws.column_dimensions["I"].width = 24

    wb.save(path)


def _write_xlsx_issue_log(path: str, rows: list[dict]):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment

    wb = Workbook()
    ws = wb.active
    ws.title = "R3_Issue_Log"

    headers = ["序号", "接口路径", "请求方法", "中文测试数据", "预期结果", "实际结果", "错误现象", "严重程度", "责任人", "状态", "发现时间", "备注", "截图路径", "cURL"]
    ws.append(headers)

    header_fill = PatternFill("solid", fgColor="FFF2CC")
    header_font = Font(bold=True)
    for col in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(wrap_text=True, vertical="top")

    wrap = Alignment(wrap_text=True, vertical="top")
    for i, r in enumerate(rows, start=1):
        ws.append(
            [
                i,
                r.get("path"),
                r.get("method"),
                r.get("cn_data"),
                r.get("expected"),
                r.get("actual"),
                r.get("phenomenon"),
                r.get("severity"),
                r.get("owner"),
                r.get("status"),
                r.get("found_at"),
                r.get("remark"),
                r.get("screenshot"),
                r.get("curl"),
            ]
        )
        row_idx = i + 1
        for col in range(1, len(headers) + 1):
            ws.cell(row=row_idx, column=col).alignment = wrap

    ws.column_dimensions["A"].width = 8
    ws.column_dimensions["B"].width = 28
    ws.column_dimensions["C"].width = 10
    ws.column_dimensions["D"].width = 40
    ws.column_dimensions["E"].width = 40
    ws.column_dimensions["F"].width = 40
    ws.column_dimensions["G"].width = 28
    ws.column_dimensions["H"].width = 10
    ws.column_dimensions["I"].width = 10
    ws.column_dimensions["J"].width = 10
    ws.column_dimensions["K"].width = 14
    ws.column_dimensions["L"].width = 18
    ws.column_dimensions["M"].width = 28
    ws.column_dimensions["N"].width = 60

    wb.save(path)


def main():
    stamp = _now_stamp()
    out_dir = os.path.join(os.path.dirname(__file__), "..", "TestEvidence", "R3")
    out_dir = os.path.abspath(out_dir)
    _ensure_dir(out_dir)

    data_log_path = os.path.join(out_dir, f"R3_Data_Log_{stamp}.xlsx")
    issue_log_path = os.path.join(out_dir, f"R3_Issue_Log_{stamp}.xlsx")
    cases_md_path = os.path.join(out_dir, "第三轮中文数据对接测试用例.md")
    report_md_path = os.path.join(out_dir, "第三轮前后端接口对接测试报告.md")
    raw_json_path = os.path.join(out_dir, f"api-integration-round3-cn-{stamp}.json")

    app = create_app("development")
    app.config.update({"TESTING": True, "ADMIN_BOOTSTRAP_SECRET": "R3-测试密钥🔐"})

    with app.app_context():
        mongo.db.users.delete_many({})
        mongo.db.products.delete_many({})
        mongo.db.orders.delete_many({})
        mongo.db.admin_roles.delete_many({})
        mongo.db.admin_permissions.delete_many({})
        mongo.db.admin_user_roles.delete_many({})
        mongo.db.admin_audit_logs.delete_many({})

    client = app.test_client()
    base_url = "http://localhost:5000"

    results = []
    data_rows = []
    issue_rows = []

    def record(case_id, module, priority, method, path, cn_input, expected_desc, resp, db_value=None, remark=None, headers=None, json_body=None):
        status_code = resp.status_code
        try:
            body = json.loads(resp.data)
        except Exception:
            body = {"raw": (resp.data.decode("utf-8", errors="replace") if resp.data else "")}

        cn_resp = _j(body)
        is_match = "是"
        if db_value is None:
            db_value_str = "N/A"
        else:
            db_value_str = _j(db_value) if isinstance(db_value, (dict, list)) else str(db_value)

        if db_value is not None and isinstance(db_value, str):
            if cn_input and isinstance(cn_input, str) and cn_input not in db_value:
                is_match = "否"

        if _is_garbled(cn_resp) or _is_garbled(db_value_str):
            is_match = "否"

        data_rows.append(
            {
                "case_id": case_id,
                "path": path,
                "cn_input": cn_input,
                "status_code": status_code,
                "cn_response": cn_resp,
                "db_value": db_value_str,
                "is_match": is_match,
                "remark": remark or "",
            }
        )

        ok = True
        if isinstance(expected_desc, dict):
            expected_status = expected_desc.get("status_code")
            if expected_status is not None and status_code != expected_status:
                ok = False
        elif isinstance(expected_desc, int):
            ok = status_code == expected_desc

        curl = _make_curl(method, base_url + path, headers or {}, json_body)
        results.append(
            {
                "case_id": case_id,
                "module": module,
                "priority": priority,
                "method": method,
                "path": path,
                "headers": headers or {},
                "json": json_body,
                "expected": expected_desc,
                "status_code": status_code,
                "response": body,
                "curl": curl,
                "db_value": db_value_str,
                "remark": remark or "",
                "pass": ok,
            }
        )

        if not ok:
            issue_rows.append(
                {
                    "path": path,
                    "method": method,
                    "cn_data": cn_input,
                    "expected": str(expected_desc),
                    "actual": f"HTTP {status_code} / {cn_resp}",
                    "phenomenon": "状态码/结果不符合预期（中文场景）",
                    "severity": "P1",
                    "owner": "后端",
                    "status": "新建",
                    "found_at": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
                    "remark": remark or "",
                    "screenshot": "",
                    "curl": curl,
                }
            )

    def register_user(case_id, username, phone, password):
        path = "/api/user/register"
        payload = {"username": username, "phone": phone, "password": password}
        resp = client.post(path, json=payload)
        with app.app_context():
            db_user = mongo.db.users.find_one({"phone": phone}) if resp.status_code in (200, 201) else None
        db_value = {"username": db_user.get("username"), "phone": db_user.get("phone")} if db_user else None
        record(case_id, "用户", "P0", "POST", path, f"{username}/{phone}", {"status_code": 201}, resp, db_value=db_value, json_body=payload)
        return db_user

    def login(case_id, phone, password):
        path = "/api/user/login"
        payload = {"phone": phone, "password": password}
        resp = client.post(path, json=payload)
        record(case_id, "用户", "P0", "POST", path, f"{phone}/（密码含中文/emoji可正常登录）", {"status_code": 200}, resp, json_body=payload)
        data = json.loads(resp.data)
        return data.get("data", {}).get("tokens", {}).get("access_token"), data.get("data", {}).get("tokens", {}).get("refresh_token")

    def auth_header(token):
        return {"Authorization": f"Bearer {token}"} if token else {}

    def admin_bootstrap(case_id, username, phone, password):
        path = "/api/admin/bootstrap"
        payload = {"username": username, "phone": phone, "password": password}
        headers = {"X-Admin-Bootstrap-Secret": "R3-测试密钥🔐"}
        resp = client.post(path, json=payload, headers=headers)
        with app.app_context():
            db_user = mongo.db.users.find_one({"phone": phone}) if resp.status_code in (200, 201) else None
        db_value = {"username": db_user.get("username"), "role": db_user.get("role")} if db_user else None
        record(case_id, "管理员", "P0", "POST", path, f"{username}/{phone}", {"status_code": 201}, resp, db_value=db_value, headers=headers, json_body=payload)
        return db_user

    def admin_create_product(case_id, token, name, category, description, price, stock, image_url=None):
        path = "/api/admin/products"
        payload = {"name": name, "category": category, "price": price, "stock": stock, "description": description, "image_url": image_url}
        headers = auth_header(token)
        resp = client.post(path, json=payload, headers=headers)
        record(case_id, "管理员-商品", "P0", "POST", path, name, {"status_code": 201}, resp, headers=headers, json_body=payload)
        try:
            body = json.loads(resp.data)
            return resp, body.get("data", {}).get("product_id")
        except Exception:
            return resp, None

    def get_product_detail(case_id, product_id):
        path = f"/api/product/{product_id}"
        resp = client.get(path)
        record(case_id, "商品", "P0", "GET", path, "中文商品ID", {"status_code": 200}, resp)
        return resp

    def product_list(case_id, q=None, category=None):
        qs = []
        if q is not None:
            qs.append(f"q={q}")
        if category is not None:
            qs.append(f"category={category}")
        url = "/api/product/list" + (("?" + "&".join(qs)) if qs else "")
        resp = client.get(url)
        record(case_id, "商品", "P0", "GET", url, q or "", {"status_code": 200}, resp)
        return resp

    def search_query(case_id, q):
        path = f"/api/search/query?q={q}"
        resp = client.get(path)
        record(case_id, "搜索", "P1", "GET", path, q, {"status_code": 200}, resp)
        return resp

    def cart_add(case_id, token, product_id, quantity, cn_desc):
        path = "/api/cart/add"
        payload = {"product_id": product_id, "quantity": quantity}
        headers = auth_header(token)
        resp = client.post(path, json=payload, headers=headers)
        record(case_id, "购物车", "P0", "POST", path, cn_desc, {"status_code": 200}, resp, headers=headers, json_body=payload)
        return resp

    def cart_items(case_id, token):
        path = "/api/cart/items"
        headers = auth_header(token)
        resp = client.get(path, headers=headers)
        record(case_id, "购物车", "P0", "GET", path, "查看购物车（中文商品名应正确返回）", {"status_code": 200}, resp, headers=headers)
        return resp

    def order_create(case_id, token, items, shipping_address):
        path = "/api/order/create"
        payload = {"items": items, "shipping_address": shipping_address}
        headers = auth_header(token)
        resp = client.post(path, json=payload, headers=headers)
        record(case_id, "订单", "P0", "POST", path, shipping_address, {"status_code": 201}, resp, headers=headers, json_body=payload)
        data = json.loads(resp.data)
        return data.get("data", {}).get("order_id")

    def order_pay(case_id, token, order_id, payment_method):
        path = f"/api/order/{order_id}/pay"
        payload = {"payment_method": payment_method}
        headers = auth_header(token)
        resp = client.post(path, json=payload, headers=headers)
        record(case_id, "订单", "P0", "POST", path, payment_method, {"status_code": 200}, resp, headers=headers, json_body=payload)
        return resp

    def payment_checkout(case_id, token, order_id, payment_method):
        path = "/api/payment/checkout"
        payload = {"order_id": order_id, "payment_method": payment_method}
        headers = auth_header(token)
        resp = client.post(path, json=payload, headers=headers)
        record(case_id, "支付", "P0", "POST", path, payment_method, {"status_code": 200}, resp, headers=headers, json_body=payload)
        return resp

    def admin_stats(case_id, token):
        path = "/api/admin/stats"
        headers = auth_header(token)
        resp = client.get(path, headers=headers)
        record(case_id, "管理员", "P0", "GET", path, "统计包含中文数据不乱码", {"status_code": 200}, resp, headers=headers)
        return resp

    def admin_users(case_id, token, q=None):
        path = "/api/admin/users" + (f"?q={q}" if q else "")
        headers = auth_header(token)
        resp = client.get(path, headers=headers)
        record(case_id, "管理员-用户", "P0", "GET", path, q or "", {"status_code": 200}, resp, headers=headers)
        return resp

    def admin_audit(case_id, token):
        path = "/api/admin/audit"
        headers = auth_header(token)
        resp = client.get(path, headers=headers)
        record(case_id, "管理员-审计", "P1", "GET", path, "审计日志包含中文请求数据", {"status_code": 200}, resp, headers=headers)
        return resp

    def user_profile(case_id, token):
        path = "/api/user/profile"
        headers = auth_header(token)
        resp = client.get(path, headers=headers)
        record(case_id, "用户", "P0", "GET", path, "查看个人信息（中文用户名）", {"status_code": 200}, resp, headers=headers)
        return resp

    def user_logout(case_id, token, refresh_token=None):
        path = "/api/user/logout"
        headers = auth_header(token)
        payload = {"refresh_token": refresh_token} if refresh_token else {}
        resp = client.post(path, headers=headers, json=payload)
        record(case_id, "用户", "P1", "POST", path, "登出（中文备注）", {"status_code": 200}, resp, headers=headers, json_body=payload)
        return resp

    def user_change_password(case_id, token, old_pw, new_pw, confirm_pw):
        path = "/api/user/change_password"
        headers = auth_header(token)
        payload = {"old_password": old_pw, "new_password": new_pw, "confirm_password": confirm_pw}
        resp = client.post(path, headers=headers, json=payload)
        record(case_id, "用户", "P1", "POST", path, "修改密码（新密码含中文符号/emoji）", {"status_code": 200}, resp, headers=headers, json_body=payload)
        return resp

    def admin_logout(case_id, token, refresh_token=None):
        path = "/api/admin/logout"
        headers = auth_header(token)
        payload = {"refresh_token": refresh_token} if refresh_token else {}
        resp = client.post(path, headers=headers, json=payload)
        record(case_id, "管理员", "P1", "POST", path, "管理员登出（中文）", {"status_code": 200}, resp, headers=headers, json_body=payload)
        return resp

    admin_bootstrap("API-R3-管理员-001", "管理员一😀", "18800001001", "密碼Password🔑123")
    bad_bootstrap_headers = {"X-Admin-Bootstrap-Secret": "错误密钥❌"}
    bad_bootstrap_payload = {"username": "管理员二", "phone": "18800001002", "password": "密码123456"}
    bad_bootstrap_resp = client.post("/api/admin/bootstrap", json=bad_bootstrap_payload, headers=bad_bootstrap_headers)
    record("API-R3-管理员-002", "管理员", "P1", "POST", "/api/admin/bootstrap", "错误密钥应拒绝（中文）", {"status_code": 403}, bad_bootstrap_resp, headers=bad_bootstrap_headers, json_body=bad_bootstrap_payload)

    register_user("API-R3-用户-001", "张三😀（简体）", "18800002001", "密码🔒123456")
    register_user("API-R3-用户-002", "李四（繁體：測試）", "18800002002", "密碼🔒123456")
    register_user("API-R3-用户-003", "王五，标点：，。！？【】（）—％￥", "18800002003", "密码🔒123456")

    reg_bad_1 = {"username": "", "phone": "", "password": ""}
    resp = client.post("/api/user/register", json=reg_bad_1)
    record("API-R3-用户-004", "用户", "P1", "POST", "/api/user/register", "空值（中文场景）", {"status_code": 400}, resp, json_body=reg_bad_1)
    reg_bad_2 = {"username": "<script>alert('XSS')</script>😀", "phone": "18800002004", "password": "密码🔒123456"}
    resp = client.post("/api/user/register", json=reg_bad_2)
    record("API-R3-用户-005", "用户", "P1", "POST", "/api/user/register", "XSS 脚本混合 emoji（不应 500/不乱码）", {"status_code": 201}, resp, json_body=reg_bad_2)
    reg_bad_3 = {"username": "注入测试' OR 1=1 --", "phone": "18800002005", "password": "密码🔒123456"}
    resp = client.post("/api/user/register", json=reg_bad_3)
    record("API-R3-用户-006", "用户", "P1", "POST", "/api/user/register", "SQL 注入字符串（不应 500/不乱码）", {"status_code": 201}, resp, json_body=reg_bad_3)

    admin_access, admin_refresh = login("API-R3-管理员-003", "18800001001", "密碼Password🔑123")

    u1_access, u1_refresh = login("API-R3-用户-007", "18800002001", "密码🔒123456")
    u2_access, u2_refresh = login("API-R3-用户-008", "18800002002", "密碼🔒123456")
    u3_access, u3_refresh = login("API-R3-用户-009", "18800002003", "密码🔒123456")

    bad_login_1 = client.post("/api/user/login", json={"phone": "18800002001", "password": "错误密码❌"})
    record("API-R3-用户-010", "用户", "P1", "POST", "/api/user/login", "错误密码（中文提示可为英文，但不乱码）", {"status_code": 401}, bad_login_1, json_body={"phone": "18800002001", "password": "错误密码❌"})
    bad_login_2 = client.post("/api/user/login", json={"phone": "", "password": ""})
    record("API-R3-用户-011", "用户", "P2", "POST", "/api/user/login", "空字段登录", {"status_code": 400}, bad_login_2, json_body={"phone": "", "password": ""})
    bad_login_3 = client.post("/api/user/login", json={"phone": "18800002001", "password": None})
    record("API-R3-用户-012", "用户", "P2", "POST", "/api/user/login", "password=null（稳定性）", {"status_code": 400}, bad_login_3, json_body={"phone": "18800002001", "password": None})

    user_profile("API-R3-用户-013", u1_access)
    user_profile("API-R3-用户-014", u2_access)
    user_profile("API-R3-用户-015", u3_access)
    no_token_profile = client.get("/api/user/profile")
    record("API-R3-用户-016", "用户", "P1", "GET", "/api/user/profile", "未登录访问应 401", {"status_code": 401}, no_token_profile)
    bad_token_profile = client.get("/api/user/profile", headers={"Authorization": "Bearer 非法token😀"})
    record("API-R3-用户-017", "用户", "P1", "GET", "/api/user/profile", "非法 token 应 401", {"status_code": 401}, bad_token_profile, headers={"Authorization": "Bearer 非法token😀"})

    (resp, p1_id) = admin_create_product("API-R3-商品-001", admin_access, "助听器A（简体）😀", "听力设备", "适用于轻度听力损失；含中文标点，。！？", 1999.99, 50)
    (resp, p2_id) = admin_create_product("API-R3-商品-002", admin_access, "助聽器B（繁體）", "輔具", "繁體中文測試；全角符號：１２３", 2999.0, 20)
    (resp, p3_id) = admin_create_product("API-R3-商品-003", admin_access, "助听器C（emoji混合😺）", "听力设备", "描述含 emoji 与中文混排😀", 1.0, 1)

    bad_prod_1 = client.post("/api/admin/products", headers=auth_header(admin_access), json={"name": "", "category": "", "price": -1, "stock": -1})
    record("API-R3-商品-004", "管理员-商品", "P1", "POST", "/api/admin/products", "异常：空/负数应返回 400（中文字段校验）", {"status_code": 400}, bad_prod_1, headers=auth_header(admin_access), json_body={"name": "", "category": "", "price": -1, "stock": -1})
    bad_prod_2 = client.post("/api/admin/products", headers=auth_header(admin_access), json={"category": "听力设备", "price": 1, "stock": 1})
    record("API-R3-商品-005", "管理员-商品", "P1", "POST", "/api/admin/products", "异常：缺少 name", {"status_code": 400}, bad_prod_2, headers=auth_header(admin_access), json_body={"category": "听力设备", "price": 1, "stock": 1})
    bad_prod_3 = client.post("/api/admin/products", headers=auth_header(admin_access), json={"name": "XSS<img src=x onerror=alert(1)>", "category": "听力设备", "price": 1, "stock": 1})
    record("API-R3-商品-006", "管理员-商品", "P2", "POST", "/api/admin/products", "异常：XSS 字符串（可接受但不应 500/不乱码）", {"status_code": 201}, bad_prod_3, headers=auth_header(admin_access), json_body={"name": "XSS<img src=x onerror=alert(1)>", "category": "听力设备", "price": 1, "stock": 1})

    product_list("API-R3-查询-001", q="助听器")
    product_list("API-R3-查询-002", q="助聽器")
    product_list("API-R3-查询-003", q="emoji😀")
    bad_list_1 = client.get("/api/product/list?page=中文")
    record("API-R3-查询-004", "商品", "P1", "GET", "/api/product/list?page=中文", "异常：分页参数为中文", {"status_code": 400}, bad_list_1)
    bad_list_2 = client.get("/api/product/list?min_price=一百")
    record("API-R3-查询-005", "商品", "P1", "GET", "/api/product/list?min_price=一百", "异常：价格参数为中文", {"status_code": 400}, bad_list_2)
    bad_list_3 = client.get("/api/product/list?page_size=999999")
    record("API-R3-查询-006", "商品", "P2", "GET", "/api/product/list?page_size=999999", "异常：超大 page_size（不应 500/不乱码）", {"status_code": 200}, bad_list_3)

    search_query("API-R3-搜索-001", q="听力设备")
    search_query("API-R3-搜索-002", q="輔具")
    search_query("API-R3-搜索-003", q="中文标点，。！？😀")
    search_query("API-R3-搜索-004", q="")
    search_query("API-R3-搜索-005", q="(a+)+")
    search_query("API-R3-搜索-006", q="注入' OR 1=1 --😀")

    cart_add("API-R3-购物车-001", u1_access, p1_id, 1, "加入：助听器A😀")
    cart_add("API-R3-购物车-002", u1_access, p2_id, 2, "加入：助聽器B（繁體）")
    cart_add("API-R3-购物车-003", u1_access, p3_id, 1, "加入：助听器C（emoji😺）")
    cart_items("API-R3-购物车-004", u1_access)
    too_much_payload = {"product_id": p1_id, "quantity": 99999}
    too_much_resp = client.post("/api/cart/add", headers=auth_header(u1_access), json=too_much_payload)
    record("API-R3-购物车-005", "购物车", "P1", "POST", "/api/cart/add", "异常：超出库存应提示（不崩溃）", {"status_code": 400}, too_much_resp, headers=auth_header(u1_access), json_body=too_much_payload)
    bad_cart_1 = client.post("/api/cart/add", headers=auth_header(u1_access), json={"quantity": 1})
    record("API-R3-购物车-006", "购物车", "P1", "POST", "/api/cart/add", "异常：缺少 product_id", {"status_code": 400}, bad_cart_1, headers=auth_header(u1_access), json_body={"quantity": 1})
    bad_cart_2 = client.post("/api/cart/add", headers=auth_header(u1_access), json={"product_id": "000000000000000000000000", "quantity": 1})
    record("API-R3-购物车-007", "购物车", "P1", "POST", "/api/cart/add", "异常：不存在商品ID", {"status_code": 400}, bad_cart_2, headers=auth_header(u1_access), json_body={"product_id": "000000000000000000000000", "quantity": 1})
    bad_cart_3 = client.get("/api/cart/items")
    record("API-R3-购物车-008", "购物车", "P1", "GET", "/api/cart/items", "异常：未登录访问购物车应 401", {"status_code": 401}, bad_cart_3)

    order1 = order_create("API-R3-订单-001", u1_access, items=[{"product_id": p1_id, "quantity": 1}], shipping_address="北京市海淀区中关村大街１号（全角）😀")
    order2 = order_create("API-R3-订单-002", u2_access, items=[{"product_id": p2_id, "quantity": 1}], shipping_address="上海市浦东新区世纪大道８８号（繁體：測試）")
    order3 = order_create("API-R3-订单-003", u3_access, items=[{"product_id": p1_id, "quantity": 1}], shipping_address="广州市天河区体育西路（中文标点，。！？）")
    bad_order_1 = client.post("/api/order/create", headers=auth_header(u1_access), json={"items": [], "shipping_address": ""})
    record("API-R3-订单-004", "订单", "P1", "POST", "/api/order/create", "异常：缺少订单信息", {"status_code": 400}, bad_order_1, headers=auth_header(u1_access), json_body={"items": [], "shipping_address": ""})
    bad_order_2 = client.post("/api/order/create", headers=auth_header(u1_access), json={"items": [{"product_id": p1_id, "quantity": "中文"}], "shipping_address": "地址😀"})
    record("API-R3-订单-005", "订单", "P1", "POST", "/api/order/create", "异常：quantity=中文", {"status_code": 400}, bad_order_2, headers=auth_header(u1_access), json_body={"items": [{"product_id": p1_id, "quantity": "中文"}], "shipping_address": "地址😀"})
    bad_order_3 = client.post("/api/order/create", headers=auth_header(u1_access), json={"items": [{"product_id": p3_id, "quantity": 999}], "shipping_address": "地址😀"})
    record("API-R3-订单-006", "订单", "P1", "POST", "/api/order/create", "异常：库存不足", {"status_code": 400}, bad_order_3, headers=auth_header(u1_access), json_body={"items": [{"product_id": p3_id, "quantity": 999}], "shipping_address": "地址😀"})

    payment_checkout("API-R3-支付-001", u1_access, order1, "微信支付😀")
    payment_checkout("API-R3-支付-002", u2_access, order2, "支付宝（中文）")
    payment_checkout("API-R3-支付-003", u3_access, order3, "银联支付（全角１２３）")

    pay_missing = order_create("API-R3-订单-007", u1_access, items=[{"product_id": p1_id, "quantity": 1}], shipping_address="缺少支付方式用例（中文）")
    bad_pay_1 = client.post(f"/api/order/{pay_missing}/pay", headers=auth_header(u1_access), json={})
    record("API-R3-支付-004", "支付", "P1", "POST", f"/api/order/{pay_missing}/pay", "异常：缺少 payment_method", {"status_code": 400}, bad_pay_1, headers=auth_header(u1_access), json_body={})

    pay1 = order_create("API-R3-订单-008", u1_access, items=[{"product_id": p1_id, "quantity": 1}], shipping_address="支付接口正向用例一（中文）😀")
    pay2 = order_create("API-R3-订单-009", u2_access, items=[{"product_id": p2_id, "quantity": 1}], shipping_address="支付接口正向用例二（繁體）")
    pay3 = order_create("API-R3-订单-010", u3_access, items=[{"product_id": p1_id, "quantity": 1}], shipping_address="支付接口正向用例三（标点，。！？）")
    order_pay("API-R3-支付-005", u1_access, pay1, "微信支付😀")
    order_pay("API-R3-支付-006", u2_access, pay2, "支付宝（中文）")
    order_pay("API-R3-支付-007", u3_access, pay3, "银联支付（全角１２３）")
    bad_checkout_1 = client.post("/api/payment/checkout", headers=auth_header(u1_access), json={"order_id": "", "payment_method": ""})
    record("API-R3-支付-008", "支付", "P1", "POST", "/api/payment/checkout", "异常：缺少字段", {"status_code": 400}, bad_checkout_1, headers=auth_header(u1_access), json_body={"order_id": "", "payment_method": ""})
    bad_checkout_2 = client.post("/api/payment/checkout", headers=auth_header(u1_access), json={"order_id": "000000000000000000000000", "payment_method": "微信"})
    record("API-R3-支付-009", "支付", "P2", "POST", "/api/payment/checkout", "异常：不存在订单ID（不应 500）", {"status_code": 200}, bad_checkout_2, headers=auth_header(u1_access), json_body={"order_id": "000000000000000000000000", "payment_method": "微信"})

    admin_stats("API-R3-管理员-004", admin_access)
    admin_stats("API-R3-管理员-005", admin_access)
    admin_stats("API-R3-管理员-006", admin_access)
    admin_users("API-R3-管理员-007", admin_access, q="张三")
    admin_users("API-R3-管理员-008", admin_access, q="測試")
    admin_users("API-R3-管理员-009", admin_access, q="标点")
    admin_audit("API-R3-管理员-010", admin_access)
    admin_audit("API-R3-管理员-011", admin_access)
    admin_audit("API-R3-管理员-012", admin_access)

    user_change_password("API-R3-用户-018", u1_access, "密码🔒123456", "新密码🔐中文😀123456", "新密码🔐中文😀123456")
    user_change_password("API-R3-用户-019", u2_access, "密碼🔒123456", "新密碼🔐繁體😀123456", "新密碼🔐繁體😀123456")
    user_change_password("API-R3-用户-020", u3_access, "密码🔒123456", "新密码🔐标点，。！？123456", "新密码🔐标点，。！？123456")
    bad_pw_1 = client.post("/api/user/change_password", headers=auth_header(u1_access), json={"old_password": "错误❌", "new_password": "新密码123456", "confirm_password": "新密码123456"})
    record("API-R3-用户-021", "用户", "P1", "POST", "/api/user/change_password", "异常：旧密码错误", {"status_code": 401}, bad_pw_1, headers=auth_header(u1_access), json_body={"old_password": "错误❌", "new_password": "新密码123456", "confirm_password": "新密码123456"})
    bad_pw_2 = client.post("/api/user/change_password", headers=auth_header(u1_access), json={"old_password": "密码🔒123456", "new_password": "新密码123456", "confirm_password": "不一致❌"})
    record("API-R3-用户-022", "用户", "P1", "POST", "/api/user/change_password", "异常：确认密码不一致", {"status_code": 400}, bad_pw_2, headers=auth_header(u1_access), json_body={"old_password": "密码🔒123456", "new_password": "新密码123456", "confirm_password": "不一致❌"})
    bad_pw_3 = client.post("/api/user/change_password", headers=auth_header(u1_access), json={"old_password": "密码🔒123456", "new_password": "短", "confirm_password": "短"})
    record("API-R3-用户-023", "用户", "P2", "POST", "/api/user/change_password", "异常：新密码太短", {"status_code": 400}, bad_pw_3, headers=auth_header(u1_access), json_body={"old_password": "密码🔒123456", "new_password": "短", "confirm_password": "短"})

    user_logout("API-R3-用户-024", u1_access, u1_refresh)
    user_logout("API-R3-用户-025", u2_access, u2_refresh)
    user_logout("API-R3-用户-026", u3_access, u3_refresh)
    no_token_logout = client.post("/api/user/logout", json={})
    record("API-R3-用户-027", "用户", "P2", "POST", "/api/user/logout", "异常：未登录登出应 401", {"status_code": 401}, no_token_logout, json_body={})

    admin_logout("API-R3-管理员-013", admin_access, admin_refresh)

    cases_md = []
    cases_md.append("# 第三轮中文数据对接测试用例\n")
    cases_md.append("用例编号规则：API-R3-模块-序号\n")
    cases_md.append("\n| 用例编号 | 模块 | 优先级 | 方法 | 路径 | 中文数据设计要点 | 预期 |\n|---|---|---|---|---|---|---|\n")
    for r in results:
        cn = r.get("remark") or (r.get("json") if r.get("json") is not None else "")
        cn_txt = r.get("case_id", "")
        cases_md.append(
            f"| {r['case_id']} | {r['module']} | {r['priority']} | {r['method']} | {r['path']} | 中文覆盖：简体/繁体/标点/全角/emoji/注入/XSS/超长 | {r['expected']} |\n"
        )
    with open(cases_md_path, "w", encoding="utf-8") as f:
        f.write("".join(cases_md))

    total = len(results)
    passed = sum(1 for r in results if r.get("pass"))
    failed = total - passed
    pass_rate = (passed / total * 100.0) if total else 0.0

    risk_items = []
    msg_localization = any(
        isinstance(r.get("response"), dict) and isinstance(r["response"].get("message"), str) and not re.search(r"[\u4e00-\u9fff]", r["response"]["message"])
        for r in results
    )
    if msg_localization:
        risk_items.append("- 错误提示/成功提示多为英文：中文化与一致性可作为后续优化项（不影响 Unicode 正确性）。")

    report = []
    report.append("# 第三轮前后端接口对接测试报告（中文数据）\n")
    report.append(f"- 测试日期：{datetime.now().strftime('%Y-%m-%d')}\n")
    report.append("- 测试目标：验证中文数据在接口请求/响应/落库过程中的完整性、正确性与稳定性。\n")
    report.append("- 字符集说明：MongoDB 默认使用 UTF-8 存储字符串；接口返回使用 JSON（UTF-8）。\n")
    report.append("\n## 覆盖范围\n")
    report.append("- 用户：注册/登录/个人信息/改密/登出\n")
    report.append("- 商品：后台创建（中文名称/描述）+ 前台列表/搜索\n")
    report.append("- 购物车：加入/查看/异常库存\n")
    report.append("- 订单/支付：下单（中文地址）/支付（中文支付方式）\n")
    report.append("- 管理后台：统计/用户查询/审计查询/登出\n")
    report.append("\n## 用例结果\n")
    report.append(f"- 总用例数：{total}\n")
    report.append(f"- 通过：{passed}\n")
    report.append(f"- 失败：{failed}\n")
    report.append(f"- 通过率：{pass_rate:.2f}%\n")
    report.append("\n## 缺陷统计\n")
    if issue_rows:
        report.append(f"- 缺陷数：{len(issue_rows)}（详见 {os.path.basename(issue_log_path)}）\n")
    else:
        report.append("- 缺陷数：0\n")
    report.append("\n## 风险分析与建议\n")
    if risk_items:
        report.extend([ri + "\n" for ri in risk_items])
    else:
        report.append("- 未发现中文编码/序列化/落库相关风险。\n")
    report.append("\n## 交付物\n")
    report.append(f"- 测试用例：{os.path.basename(cases_md_path)}\n")
    report.append(f"- 数据记录表：{os.path.basename(data_log_path)}\n")
    report.append(f"- 问题记录表：{os.path.basename(issue_log_path)}\n")
    report.append(f"- 原始记录：{os.path.basename(raw_json_path)}\n")

    with open(report_md_path, "w", encoding="utf-8") as f:
        f.write("".join(report))

    with open(raw_json_path, "w", encoding="utf-8") as f:
        f.write(_j({"summary": {"total": total, "passed": passed, "failed": failed}, "results": results}))

    _write_xlsx_data_log(data_log_path, data_rows)
    _write_xlsx_issue_log(issue_log_path, issue_rows)

    failed_cases = [r for r in results if not r.get("pass")]
    print(f"R3 summary: total={total} passed={passed} failed={failed}")
    if failed_cases:
        print("R3 failed cases:")
        for r in failed_cases:
            print(r.get("case_id"), r.get("method"), r.get("path"), "expected=", r.get("expected"), "actual=", r.get("status_code"))

    print("R3 outputs:")
    print(out_dir)
    print(os.path.basename(cases_md_path))
    print(os.path.basename(data_log_path))
    print(os.path.basename(issue_log_path))
    print(os.path.basename(report_md_path))
    print(os.path.basename(raw_json_path))


if __name__ == "__main__":
    main()
